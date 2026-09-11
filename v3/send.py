import smtplib
import ssl
from email.header import Header
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook


EXCEL_PATH = Path(__file__).with_name("蒋美玲.xlsx")
SENDER = "2929769973@qq.com"
EMAIL_PASSWORD = "your password"

SUBJECT = "复旦大学机考信息通知"

BODY_TEMPLATE = """{name}同学你好！我是复旦大学联络员，下面是机考信息，请仔细阅读；如有问题，请微信联系：snce147

一、考试信息

报名号：{registration_no}

机考时间：13:00-16:00（闭卷），其中 13:00-13:30 为上机环境熟悉时间，13:30 正式开始考试。

机考地点：上海市杨浦区邯郸路 220 号，复旦大学邯郸校区；{computer_room}

入校凭证：考生凭本人身份证直接刷证进校。

二、入场要求

1. 机考现场将核验考生身份证、学生证原件。

2. 机考时提交以下纸质材料：

① 《复旦大学全国优秀大学生夏令营（推免生预报名）申请表》（报考系统自动生成，自行打印），须本人签字、本科院校教务或院系在成绩排名栏“成绩排名”栏加盖公章；

② 本科阶段历年成绩单原件（须加盖教务处或院系公章）；

③ 外语水平证明复印件，如CET-4、CET-6、雅思、托福、专业外语成绩等（均应在有效期内）须出示原件验证；

④ 考生诚信考核承诺书；

⑤ 申请直接攻博生，除上述材料外还须提交拟攻读博士学位的科研计划书；

⑥ 其他支撑材料复印件，如论文、出版物等。

三、考场纪律与保密要求

1. 严禁携带智能手表、智能眼镜等电子产品进入考场。不在考核过程中借助他人或相关资料，不使用通讯、电子设备和人工智能设备。

2. 考核期间禁止拍照、录音、录像、直播，在复试期间及学校复试工作全部结束前，不以任何形式透露传播试题内容、复试考核过程等有关情况。

3. 机考结束后请立即离场，请勿在楼宇内逗留。
"""


def read_students():
    workbook = load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]

    required = ("姓名", "报名号", "机房")
    missing = [column for column in required if column not in headers]
    if missing:
        raise ValueError(f"Excel 缺少必要列：{', '.join(missing)}")

    indexes = {column: headers.index(column) for column in required}
    email_column = next(
        (column for column in ("电子邮箱", "电子信箱") if column in headers), None
    )
    if email_column is None:
        raise ValueError("Excel 缺少必要列：电子邮箱")
    indexes["电子邮箱"] = headers.index(email_column)
    students = []
    for row in rows:
        name = row[indexes["姓名"]]
        registration_no = row[indexes["报名号"]]
        computer_room = row[indexes["机房"]]
        email = row[indexes["电子邮箱"]]
        if name and registration_no and computer_room and email:
            students.append(
                {
                    "name": str(name).strip(),
                    "registration_no": str(registration_no).strip(),
                    "computer_room": str(computer_room).strip(),
                    "email": str(email).strip(),
                }
            )
    return students


def build_message(student):
    message = EmailMessage()
    message["From"] = SENDER
    message["To"] = student["email"]
    message["Subject"] = Header(
        f"{SUBJECT} - {student['name']}同学", "utf-8"
    ).encode()
    message.set_content(BODY_TEMPLATE.format(**student))
    return message


def main():
    if EMAIL_PASSWORD == "请在这里填写QQ邮箱授权码":
        raise ValueError("请先在 EMAIL_PASSWORD 中填写 QQ 邮箱授权码，不是 QQ 登录密码。")

    students = read_students()
    if not students:
        raise ValueError("Excel 中没有可发送的有效记录。")

    print(f"本次将发送 {len(students)} 封邮件：")
    for number, student in enumerate(students, 1):
        print(
            f"{number}. {student['name']}，报名号 {student['registration_no']}，"
            f"机房 {student['computer_room']}，电子邮箱 {student['email']}"
        )

    if input("确认发送请输入 SEND：").strip() != "SEND":
        print("已取消，未发送任何邮件。")
        return

    context = ssl.create_default_context()
    with smtplib.SMTP_SSL("smtp.qq.com", 465, context=context) as smtp:
        smtp.login(SENDER, EMAIL_PASSWORD)
        success_count = 0
        failure_count = 0
        for student in students:
            try:
                smtp.send_message(build_message(student))
            except Exception as error:
                failure_count += 1
                print(f"发送失败：{student['name']} -> {student['email']}，原因：{error}")
                continue

            success_count += 1
            print(f"已发送：{student['name']} -> {student['email']}")

    print(f"发送完成：成功 {success_count} 封，失败 {failure_count} 封。")


if __name__ == "__main__":
    main()
